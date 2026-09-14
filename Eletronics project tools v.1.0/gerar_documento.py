"""
gerar_documento.py
-------------------
Gera a proposta de orçamento em Word (.docx), reproduzindo o "Resumo preliminar
do orçamento" e a "Composição comercial" vistos na aba Calc_Eletrico_BT da planilha.
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
CINZA = RGBColor(0x59, 0x59, 0x59)


def _shade_cell(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_col_widths(table, widths_cm):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _titulo(doc, texto, tamanho=16, cor=AZUL):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(tamanho)
    run.font.color.rgb = cor
    return p


def _subtitulo(doc, texto):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = AZUL
    p.space_before = Pt(10)
    return p


def _tabela_chave_valor(doc, linhas, header=None, widths=(7.0, 8.5)):
    n = len(linhas) + (1 if header else 0)
    table = doc.add_table(rows=n, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(table, widths)

    r0 = 0
    if header:
        for c, txt in enumerate(header):
            cell = table.rows[0].cells[c]
            cell.text = txt
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade_cell(cell, "1F3A5F")
        r0 = 1

    for i, (k, v) in enumerate(linhas):
        row = table.rows[i + r0]
        row.cells[0].text = str(k)
        row.cells[1].text = str(v)
        row.cells[0].paragraphs[0].runs[0].bold = True
    return table


def _fmt_moeda(v):
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_pct(v):
    return f"{v * 100:.1f}%".replace(".", ",")


def gerar_proposta(cliente: dict, resultado: dict, nivel_entrega: str, prazo: str,
                    caminho_saida: str, numero_proposta: str = None) -> str:
    """
    cliente:  dicionário com os dados da aba 'Cadastro do cliente'
    resultado: dicionário retornado por calculo.calcular_orcamento()
    caminho_saida: caminho completo do arquivo .docx a ser gerado
    """
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    if not numero_proposta:
        numero_proposta = datetime.now().strftime("%Y%m%d-%H%M")

    # ---------- Cabeçalho ----------
    _titulo(doc, "PROPOSTA DE ORÇAMENTO DE PROJETO", 18)
    p = doc.add_paragraph()
    run = p.add_run(f"Proposta nº {numero_proposta}  •  Emitida em {datetime.now().strftime('%d/%m/%Y')}")
    run.font.size = Pt(10)
    run.font.color.rgb = CINZA
    doc.add_paragraph()

    # ---------- Dados do cliente ----------
    _subtitulo(doc, "1. Dados do Cliente / Obra")
    linhas_cliente = [
        ("Cliente / Empresa", cliente.get("cliente_empresa", "")),
        ("Endereço da obra", cliente.get("endereco_obra", "")),
        ("Cidade/UF", cliente.get("cidade_uf", "")),
        ("Tipo de edificação", cliente.get("tipo_edificacao", "")),
        ("Uso da edificação", cliente.get("uso_edificacao", "")),
        ("Área aproximada (m²)", cliente.get("area_m2", "")),
        ("Obra nova ou existente", cliente.get("situacao_obra", "")),
        ("Nível de entrega", nivel_entrega),
        ("Prazo", prazo),
    ]
    if cliente.get("observacoes"):
        linhas_cliente.append(("Observações", cliente.get("observacoes")))
    _tabela_chave_valor(doc, linhas_cliente)
    doc.add_paragraph()

    # ---------- Serviços orçados (resumo técnico) ----------
    _subtitulo(doc, "2. Resumo Técnico por Serviço")
    servicos = resultado["servicos"]
    table = doc.add_table(rows=1 + len(servicos), cols=6)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(table, [3.5, 2.8, 2.0, 2.2, 2.0, 2.5])

    cab = ["Serviço", "Indicador principal", "Quantidade", "Valor unit. base", "Fator total", "Valor técnico"]
    for c, txt in enumerate(cab):
        cell = table.rows[0].cells[c]
        cell.text = txt
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        _shade_cell(cell, "1F3A5F")

    for i, s in enumerate(servicos):
        row = table.rows[i + 1]
        valores = [
            s.tipo_projeto,
            s.indicador_principal,
            f"{s.quantidade:g}",
            _fmt_moeda(s.valor_unitario_base),
            f"{s.fator_total:.3f}",
            _fmt_moeda(s.valor_tecnico),
        ]
        for c, txt in enumerate(valores):
            row.cells[c].text = str(txt)
            row.cells[c].paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph()

    # ---------- Composição comercial ----------
    _subtitulo(doc, "3. Composição Comercial")
    linhas_comercial = [
        ("Subtotal técnico (serviços)", _fmt_moeda(resultado["subtotal_tecnico"])),
    ]
    if resultado["adicionais"]:
        for ad in resultado["adicionais"]:
            linhas_comercial.append((ad["nome"], _fmt_moeda(ad["valor"])))
    linhas_comercial += [
        ("Custos diretos (adicionais)", _fmt_moeda(resultado["custos_diretos"])),
        (f"Margem ({_fmt_pct(resultado['margem_pct'])})", _fmt_moeda(resultado["valor_margem"])),
        (f"Desconto ({_fmt_pct(resultado['desconto_pct'])})", "- " + _fmt_moeda(resultado["valor_desconto"])),
        (f"Imposto / NF ({_fmt_pct(resultado['imposto_pct'])})", _fmt_moeda(resultado["valor_imposto"])),
    ]
    _tabela_chave_valor(doc, linhas_comercial)
    doc.add_paragraph()

    # ---------- Valor final em destaque ----------
    table_final = doc.add_table(rows=1, cols=2)
    table_final.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(table_final, [7.0, 8.5])
    c0, c1 = table_final.rows[0].cells
    c0.text = "VALOR FINAL ESTIMADO"
    c1.text = _fmt_moeda(resultado["valor_final"])
    for c in (c0, c1):
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(13)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(c, "1F3A5F")
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "Este documento apresenta uma estimativa preliminar de orçamento, calculada com base nos "
        "indicadores e fatores técnicos cadastrados. Valores finais podem ser ajustados após "
        "levantamento detalhado do projeto."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = CINZA

    os.makedirs(os.path.dirname(caminho_saida), exist_ok=True)
    doc.save(caminho_saida)
    return caminho_saida
